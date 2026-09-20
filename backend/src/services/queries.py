"""Business-level Query operations backed exclusively by v3 ``access`` objects."""

from __future__ import annotations

import csv
import zipfile
from collections.abc import Callable
from dataclasses import asdict, dataclass, field
from io import TextIOWrapper
from pathlib import Path
from time import perf_counter

from openpyxl import Workbook

from access.cache import CacheStore
from access.catalog import CatalogRepository
from access.database import DryDataDatabase
from access.export import AndExport, ExportCondition, ExportRepository, OrExport
from access.imports import ImportRepository
from access.molecules import EntryResultFilter, MoleculeRepository
from access.packages import PackageDescriptor, PackageStore
from access.paths import DryDataPaths
from access.previews import PreviewRepository
from access.search import SearchCondition as AccessSearchCondition
from access.search import SearchRepository
from access.source_files import SourceFileStore
from core.config import Settings
from schemas.builder import ColumnProfile, SourceFileProfile, SourceFileSummary
from schemas.catalog import (
    AttributeRecord,
    AttributeStatistics,
    DatabaseStats,
    EntryIndexRecord,
    EntryRecord,
    ImportRecord,
    Molecule3DStructure,
    MoleculeAttributeAnnotation,
    MoleculeDetail,
    MoleculeListResult,
    MoleculeSummary,
    SearchExportResult,
    SearchExportSheet,
    SearchRequest,
    SearchResult,
    SmilesExportResult,
    StatisticsPreheatResult,
)
from schemas.packages import PackageDetail, PackageSummary
from services.structures import mol_block_from_smiles

_SUMMARY_HEADER = ("#", "Attribute", "Entry", "Operator", "Value", "Unit", "Matched Molecules")
_INVALID_SHEET_TITLE_CHARACTERS = set(":\\/?*[]")


@dataclass(frozen=True, slots=True)
class DryDataQueries:
    """Coordinate reads without owning SQL or filesystem paths."""

    settings: Settings
    paths: DryDataPaths = field(init=False)
    database: DryDataDatabase = field(init=False)
    molecules: MoleculeRepository = field(init=False)
    catalog: CatalogRepository = field(init=False)
    imports: ImportRepository = field(init=False)
    packages: PackageStore = field(init=False)
    previews: PreviewRepository = field(init=False)
    search_repository: SearchRepository = field(init=False)
    export_repository: ExportRepository = field(init=False)
    source_files: SourceFileStore = field(init=False)

    def __post_init__(self) -> None:
        paths = DryDataPaths(self.settings.data_root)
        database = DryDataDatabase(
            paths,
            config={
                "memory_limit": self.settings.memory_limit,
                "threads": str(self.settings.threads),
            },
        )
        packages = PackageStore(paths)
        object.__setattr__(self, "paths", paths)
        object.__setattr__(self, "database", database)
        object.__setattr__(self, "molecules", MoleculeRepository(database))
        object.__setattr__(self, "catalog", CatalogRepository(database, CacheStore(paths)))
        object.__setattr__(self, "imports", ImportRepository(database))
        object.__setattr__(self, "packages", packages)
        object.__setattr__(self, "previews", PreviewRepository(database, packages))
        object.__setattr__(self, "search_repository", SearchRepository(database))
        object.__setattr__(self, "export_repository", ExportRepository(database))
        object.__setattr__(self, "source_files", SourceFileStore(paths))

    def list_source_files(self) -> list[SourceFileSummary]:
        return [SourceFileSummary(**asdict(item)) for item in self.source_files.list_files()]

    def profile_source_file(self, name: str) -> SourceFileProfile:
        profile = self.source_files.profile(name)
        return SourceFileProfile(
            name=profile.name,
            encoding=profile.encoding,
            delimiter=profile.delimiter,
            columns=[ColumnProfile(**asdict(column)) for column in profile.columns],
            total_rows=profile.total_rows,
            preview_rows=[list(row) for row in profile.preview_rows],
        )

    def list_imports(self, *, limit: int, offset: int) -> list[ImportRecord]:
        return [
            ImportRecord(**asdict(item)) for item in self.imports.list(limit=limit, offset=offset)
        ]

    def get_import(self, import_id: int) -> ImportRecord:
        item = self.imports.get(import_id)
        if item is None:
            raise KeyError(import_id)
        return ImportRecord(**asdict(item))

    def list_molecules(
        self,
        *,
        query: str | None,
        limit: int,
        offset: int,
        entry_filter: EntryResultFilter | None = None,
    ) -> MoleculeListResult:
        page = self.molecules.list_page(
            query=query, limit=limit, offset=offset, entry_filter=entry_filter
        )
        return MoleculeListResult(
            items=[MoleculeSummary(**asdict(item)) for item in page.records],
            total=page.total,
        )

    def list_entries_index(self) -> list[EntryIndexRecord]:
        return [EntryIndexRecord(**asdict(item)) for item in self.catalog.list_entries_index()]

    def get_molecule(self, molecule_id: int) -> MoleculeDetail:
        molecule = self.molecules.get(molecule_id)
        if molecule is None:
            raise KeyError(molecule_id)
        return MoleculeDetail(**asdict(molecule))

    def get_molecule_3d_structure(self, molecule_id: int) -> Molecule3DStructure:
        """Read one molecule then derive a non-persisted 3D structure from its SMILES."""
        molecule = self.molecules.get(molecule_id)
        if molecule is None:
            raise KeyError(molecule_id)
        return Molecule3DStructure(
            molecule_id=molecule.molecule_id,
            canonical_smiles=molecule.canonical_smiles,
            mol_block=mol_block_from_smiles(molecule.canonical_smiles),
        )

    def list_molecule_attributes(self, molecule_id: int) -> list[AttributeRecord]:
        return [
            AttributeRecord(**asdict(item))
            for item in self.catalog.molecule_attributes(molecule_id)
        ]

    def list_molecule_attribute_annotations(
        self, molecule_id: int, attribute_id: int
    ) -> list[MoleculeAttributeAnnotation]:
        return [
            MoleculeAttributeAnnotation(**asdict(item))
            for item in self.catalog.molecule_attribute_annotations(molecule_id, attribute_id)
        ]

    def list_attributes(self) -> list[AttributeRecord]:
        return [AttributeRecord(**asdict(item)) for item in self.catalog.list_attributes()]

    def list_entries(self, attribute_id: int) -> list[EntryRecord]:
        return [EntryRecord(**asdict(item)) for item in self.catalog.list_entries(attribute_id)]

    def attribute_statistics(self, attribute_id: int) -> list[AttributeStatistics]:
        return [
            AttributeStatistics(**asdict(item))
            for item in self.catalog.attribute_statistics(attribute_id)
        ]

    def preheat_attribute_statistics(
        self, progress: Callable[[int, int, str], None] | None = None
    ) -> StatisticsPreheatResult:
        """Compute and cache statistics for every Attribute so later reads are instant."""
        started = perf_counter()
        attributes = self.list_attributes()
        total = len(attributes)
        preheated = 0
        failed = 0
        for index, attribute in enumerate(attributes, start=1):
            try:
                self.catalog.attribute_statistics(attribute.attribute_id)
            except Exception:
                failed += 1
            else:
                preheated += 1
            if progress is not None:
                progress(index, total, attribute.attribute_name)
        return StatisticsPreheatResult(
            total=total,
            preheated=preheated,
            failed=failed,
            elapsed_seconds=perf_counter() - started,
        )

    def database_statistics(self, *, jobs: int) -> DatabaseStats:
        return DatabaseStats(**self.catalog.counts(), jobs=jobs)

    def search(self, request: SearchRequest) -> SearchResult:
        started = perf_counter()
        conditions = [
            AccessSearchCondition(**condition.model_dump()) for condition in request.conditions
        ]
        page = self.search_repository.search(
            conditions,
            logic=request.logic,
            limit=request.limit,
            offset=request.offset,
        )
        return SearchResult(
            molecules=[MoleculeSummary(**asdict(item)) for item in page.records],
            elapsed_ms=(perf_counter() - started) * 1000,
            limit=request.limit,
            offset=request.offset,
            total=page.total,
        )

    def export_search_workbook(
        self,
        request: SearchRequest,
        *,
        file_stem: str,
    ) -> SearchExportResult:
        """Read one full search result set and write its XLSX export workbook."""
        conditions = [
            AccessSearchCondition(**condition.model_dump()) for condition in request.conditions
        ]
        export = self.export_repository.export(conditions, logic=request.logic)
        return _write_search_export_workbook(
            self.paths.exports_directory,
            file_stem=file_stem,
            export=export,
        )

    def export_smiles_archive(
        self,
        *,
        rows_per_file: int,
        file_stem: str,
        entry_filter: EntryResultFilter | None = None,
    ) -> SmilesExportResult:
        """Write molecule SMILES into chunked single-column CSV files packed as one ZIP."""
        smiles = self.molecules.list_smiles(entry_filter=entry_filter)
        return _write_smiles_export_archive(
            self.paths.exports_directory,
            file_stem=file_stem,
            smiles=smiles,
            rows_per_file=rows_per_file,
        )

    def list_packages(self) -> list[PackageSummary]:
        return [self._package_summary(item) for item in self.packages.list_incoming()]

    def get_package(self, package_id: str) -> PackageDetail:
        item = self.packages.descriptor(package_id)
        return PackageDetail(**self._package_summary(item).model_dump(), files=list(item.files))

    def preview_package(self, package_id: str):
        item = self.packages.descriptor(package_id)
        if item.kind == "molecules":
            return self.previews.molecule_package(package_id)
        return self.previews.annotation_package(package_id)

    @staticmethod
    def _package_summary(item: PackageDescriptor) -> PackageSummary:
        return PackageSummary(
            package_id=item.package_id,
            package_kind=item.kind,
            package_name=item.name,
            size_bytes=item.size_bytes,
            modified_at=item.modified_at,
        )


def _write_search_export_workbook(
    directory: Path,
    *,
    file_stem: str,
    export: AndExport | OrExport,
) -> SearchExportResult:
    directory.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    sheets = [_append_summary_sheet(workbook, export)]
    if isinstance(export, AndExport):
        sheets.append(_append_and_sheet(workbook, export))
    else:
        sheets.extend(_append_or_sheets(workbook, export))
    file_name = f"{file_stem}.xlsx"
    workbook.save(directory / file_name)
    return SearchExportResult(
        file_name=file_name,
        logic="and" if isinstance(export, AndExport) else "or",
        total=export.total,
        sheets=sheets,
    )


def _write_smiles_export_archive(
    directory: Path,
    *,
    file_stem: str,
    smiles: list[str],
    rows_per_file: int,
) -> SmilesExportResult:
    if not smiles:
        raise ValueError("The molecule registry is empty; there is nothing to export.")
    directory.mkdir(parents=True, exist_ok=True)
    file_name = f"{file_stem}.zip"
    files: list[str] = []
    with zipfile.ZipFile(directory / file_name, "w", zipfile.ZIP_DEFLATED) as archive:
        for part_index, start in enumerate(range(0, len(smiles), rows_per_file), start=1):
            part_name = f"smiles_{part_index:04d}.csv"
            with archive.open(part_name, "w") as raw:
                with TextIOWrapper(raw, encoding="utf-8", newline="") as stream:
                    writer = csv.writer(stream)
                    writer.writerow(["SMILES"])
                    writer.writerows((value,) for value in smiles[start : start + rows_per_file])
            files.append(part_name)
    return SmilesExportResult(
        file_name=file_name,
        total=len(smiles),
        rows_per_file=rows_per_file,
        files=files,
    )


def _append_summary_sheet(
    workbook: Workbook,
    export: AndExport | OrExport,
) -> SearchExportSheet:
    sheet = workbook.create_sheet("Summary")
    sheet.append(_SUMMARY_HEADER)
    for index, condition in enumerate(export.conditions, start=1):
        sheet.append(
            [
                index,
                condition.attribute_name,
                condition.entry_key,
                condition.operator,
                _condition_value_text(condition),
                condition.unit,
                condition.matched_count,
            ]
        )
    scope = (
        "Intersection of all conditions"
        if isinstance(export, AndExport)
        else "Union of all conditions (distinct molecules)"
    )
    sheet.append(["Total", scope, None, None, None, None, export.total])
    return SearchExportSheet(title="Summary", row_count=len(export.conditions) + 1)


def _append_and_sheet(workbook: Workbook, export: AndExport) -> SearchExportSheet:
    sheet = workbook.create_sheet("Molecules")
    sheet.append(["Lab ID", "SMILES", *_value_headers(export.conditions)])
    rows = 0
    for molecule, values in zip(export.rows.molecules, export.rows.values):
        sheet.append([molecule.lab_id, molecule.canonical_smiles, *values])
        rows += 1
    return SearchExportSheet(title="Molecules", row_count=rows)


def _append_or_sheets(workbook: Workbook, export: OrExport) -> list[SearchExportSheet]:
    sheets: list[SearchExportSheet] = []
    for position, item in enumerate(export.sets, start=1):
        condition = export.conditions[item.condition_index]
        title = _sheet_title(f"Condition {position} - {condition.attribute_name}")
        sheet = workbook.create_sheet(title)
        sheet.append(["Lab ID", "SMILES", _value_header(position, condition)])
        rows = 0
        for molecule, value in zip(item.molecules, item.values):
            sheet.append([molecule.lab_id, molecule.canonical_smiles, value])
            rows += 1
        sheets.append(SearchExportSheet(title=title, row_count=rows))
    return sheets


def _value_headers(conditions: tuple[ExportCondition, ...]) -> list[str]:
    return [
        _value_header(position, condition)
        for position, condition in enumerate(conditions, start=1)
    ]


def _value_header(position: int, condition: ExportCondition) -> str:
    label = f"Condition {position} - {condition.attribute_name}"
    if condition.unit:
        return f"{label} ({condition.unit})"
    return label


def _sheet_title(raw: str) -> str:
    cleaned = "".join(
        "-" if character in _INVALID_SHEET_TITLE_CHARACTERS else character for character in raw
    )
    return cleaned[:31]


def _condition_value_text(condition: ExportCondition) -> str:
    if condition.second_value is not None:
        return f"{condition.value} to {condition.second_value}"
    return str(condition.value)
