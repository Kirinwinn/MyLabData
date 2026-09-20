"""Integration tests for the search export Job and its XLSX workbook download."""

from io import BytesIO
from pathlib import Path
from time import monotonic, sleep

import duckdb
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core.config import Settings
from main import create_app

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _create_v3_database(root: Path) -> None:
    root.mkdir(parents=True, exist_ok=True)
    with duckdb.connect(str(root / "DryData.duckdb")) as database:
        database.execute(
            """
            CREATE TABLE Molecules (molecule_id BIGINT, lab_id VARCHAR,
              canonical_smiles VARCHAR, channel VARCHAR, created_at TIMESTAMP);
            CREATE TABLE Attributes (attribute_id BIGINT, attribute_key VARCHAR,
              attribute_name VARCHAR, value_type VARCHAR, unit VARCHAR,
              description VARCHAR);
            CREATE TABLE Entries (entry_id BIGINT, attribute_id BIGINT,
              entry_key VARCHAR, annotation_kind VARCHAR, method_name VARCHAR,
              method_version VARCHAR, conditions_json JSON, is_mutable BOOLEAN,
              description VARCHAR);
            CREATE TABLE Annotations (molecule_id BIGINT, entry_id BIGINT,
              value_number DOUBLE, value_text VARCHAR, value_boolean BOOLEAN,
              created_at TIMESTAMP, updated_at TIMESTAMP);
            CREATE TABLE Sources (source_id BIGINT, source_key VARCHAR,
              source_type VARCHAR, description VARCHAR);
            CREATE TABLE Imports (import_id BIGINT, source_id BIGINT,
              file_hash VARCHAR, data_type VARCHAR, status VARCHAR,
              processor VARCHAR, total_rows BIGINT, error_message VARCHAR,
              created_at TIMESTAMP, finished_at TIMESTAMP);
            """
        )
        database.execute(
            "INSERT INTO Attributes VALUES (1, 'emission', 'Emission Wavelength', 'number',"
            " 'nm', NULL), (2, 'note', 'Note', 'text', NULL, NULL)"
        )
        database.execute(
            "INSERT INTO Entries VALUES (1, 1, 'emission.proby.CS(C)=O', 'calculation', 'T',"
            " '1', '{}', FALSE, NULL), (2, 2, 'note.manual', 'manual', 'T', '1', '{}', TRUE, NULL)"
        )
        database.execute(
            "INSERT INTO Molecules VALUES (10, 'L10', 'CCO', NULL, '2026-09-01'),"
            " (11, 'L11', 'CCC', NULL, '2026-09-02'), (12, 'L12', 'CCN', NULL, '2026-09-03')"
        )
        database.execute(
            "INSERT INTO Annotations VALUES (10, 1, 700, NULL, NULL, '2026-09-01', '2026-09-01'),"
            " (11, 1, 800, NULL, NULL, '2026-09-02', '2026-09-02'),"
            " (12, 1, 700, NULL, NULL, '2026-09-03', '2026-09-03'),"
            " (10, 2, NULL, 'blue', NULL, '2026-09-01', '2026-09-01'),"
            " (11, 2, NULL, 'red', NULL, '2026-09-02', '2026-09-02')"
        )


def _wait_for_job(client: TestClient, job_id: str, *, timeout: float = 10) -> dict:
    deadline = monotonic() + timeout
    while monotonic() < deadline:
        job = client.get(f"/api/v1/jobs/{job_id}").json()
        if job["status"] in {"completed", "failed", "cancelled"}:
            return job
        sleep(0.02)
    raise AssertionError(f"Job {job_id} did not finish")


def _export_job(client: TestClient, payload: dict) -> dict:
    accepted = client.post("/api/v1/search/export", json=payload)
    assert accepted.status_code == 202, accepted.text
    job = _wait_for_job(client, accepted.json()["job_id"])
    assert job["status"] == "completed", job
    return job


def test_and_export_writes_summary_and_molecules_sheets(tmp_path: Path) -> None:
    settings = Settings(data_root=tmp_path / "DryData", memory_limit="1GB", threads=1)
    _create_v3_database(settings.data_root)
    application = create_app(settings)

    with TestClient(application) as client:
        job = _export_job(
            client,
            {
                "conditions": [
                    {"attribute_id": 1, "entry_id": 1, "operator": "gt", "value": 650},
                    {"attribute_id": 2, "entry_id": 2, "operator": "contains", "value": "r"},
                ],
                "logic": "and",
            },
        )

        result = job["result"]
        assert result["logic"] == "and"
        assert result["total"] == 1
        assert [(sheet["title"], sheet["row_count"]) for sheet in result["sheets"]] == [
            ("Summary", 3),
            ("Molecules", 1),
        ]
        workbook_path = settings.data_root / "Temp" / "exports" / result["file_name"]
        assert workbook_path.is_file()

        download = client.get(f"/api/v1/jobs/{job['job_id']}/export")
        assert download.status_code == 200
        assert download.headers["content-type"] == XLSX_MEDIA_TYPE
        book = load_workbook(BytesIO(download.content))
        assert book.sheetnames == ["Summary", "Molecules"]

        summary = list(book["Summary"].iter_rows(values_only=True))
        assert summary[0] == (
            "#",
            "Attribute",
            "Entry",
            "Operator",
            "Value",
            "Unit",
            "Matched Molecules",
        )
        assert summary[1][:6] == (
            1,
            "Emission Wavelength",
            "emission.proby.CS(C)=O",
            "gt",
            "650",
            "nm",
        )
        assert summary[1][6] == 3
        assert summary[2][:6] == (2, "Note", "note.manual", "contains", "r", None)
        assert summary[2][6] == 1
        assert summary[3] == (
            "Total",
            "Intersection of all conditions",
            None,
            None,
            None,
            None,
            1,
        )

        molecules = list(book["Molecules"].iter_rows(values_only=True))
        assert molecules[0] == (
            "Lab ID",
            "SMILES",
            "Condition 1 - Emission Wavelength (nm)",
            "Condition 2 - Note",
        )
        assert len(molecules) == 2
        assert molecules[1][0] == "L11"
        assert molecules[1][1] == "CCC"
        assert molecules[1][2] == 800
        assert molecules[1][3] == "red"


def test_or_export_writes_one_sheet_per_condition(tmp_path: Path) -> None:
    settings = Settings(data_root=tmp_path / "DryData", memory_limit="1GB", threads=1)
    _create_v3_database(settings.data_root)
    application = create_app(settings)

    with TestClient(application) as client:
        job = _export_job(
            client,
            {
                "conditions": [
                    {
                        "attribute_id": 1,
                        "entry_id": 1,
                        "operator": "between",
                        "value": 650,
                        "second_value": 750,
                    },
                    {"attribute_id": 2, "entry_id": 2, "operator": "contains", "value": "r"},
                ],
                "logic": "or",
            },
        )

        result = job["result"]
        assert result["logic"] == "or"
        assert result["total"] == 3
        assert [(sheet["title"], sheet["row_count"]) for sheet in result["sheets"]] == [
            ("Summary", 3),
            ("Condition 1 - Emission Waveleng", 2),
            ("Condition 2 - Note", 1),
        ]

        book = load_workbook(settings.data_root / "Temp" / "exports" / result["file_name"])
        assert book.sheetnames == [
            "Summary",
            "Condition 1 - Emission Waveleng",
            "Condition 2 - Note",
        ]

        summary = list(book["Summary"].iter_rows(values_only=True))
        assert summary[1][:6] == (
            1,
            "Emission Wavelength",
            "emission.proby.CS(C)=O",
            "between",
            "650 to 750",
            "nm",
        )
        assert summary[1][6] == 2
        assert summary[2][6] == 1
        assert summary[3] == (
            "Total",
            "Union of all conditions (distinct molecules)",
            None,
            None,
            None,
            None,
            3,
        )

        first = list(book["Condition 1 - Emission Waveleng"].iter_rows(values_only=True))
        assert first[0] == ("Lab ID", "SMILES", "Condition 1 - Emission Wavelength (nm)")
        assert [row[0] for row in first[1:]] == ["L10", "L12"]
        assert [row[1] for row in first[1:]] == ["CCO", "CCN"]
        assert [row[2] for row in first[1:]] == [700, 700]

        second = list(book["Condition 2 - Note"].iter_rows(values_only=True))
        assert second[0] == ("Lab ID", "SMILES", "Condition 2 - Note")
        assert second[1] == ("L11", "CCC", "red")


def test_export_download_rejects_missing_or_unrelated_jobs(tmp_path: Path) -> None:
    settings = Settings(data_root=tmp_path / "DryData", memory_limit="1GB", threads=1)
    _create_v3_database(settings.data_root)
    application = create_app(settings)

    with TestClient(application) as client:
        assert client.get("/api/v1/jobs/does-not-exist/export").status_code == 404

        search_job = _wait_for_job(
            client,
            client.post(
                "/api/v1/search",
                json={
                    "conditions": [
                        {"attribute_id": 1, "entry_id": 1, "operator": "gt", "value": 650}
                    ],
                    "logic": "and",
                },
            ).json()["job_id"],
        )
        assert search_job["status"] == "completed"
        assert client.get(f"/api/v1/jobs/{search_job['job_id']}/export").status_code == 404

        job = _export_job(
            client,
            {
                "conditions": [
                    {"attribute_id": 1, "entry_id": 1, "operator": "gt", "value": 650}
                ],
                "logic": "and",
            },
        )
        file_name = job["result"]["file_name"]
        assert client.get(f"/api/v1/jobs/{job['job_id']}/export").status_code == 200
        (settings.data_root / "Temp" / "exports" / file_name).unlink()
        assert client.get(f"/api/v1/jobs/{job['job_id']}/export").status_code == 404


def test_export_request_validates_payload(tmp_path: Path) -> None:
    settings = Settings(data_root=tmp_path / "DryData", memory_limit="1GB", threads=1)
    _create_v3_database(settings.data_root)
    application = create_app(settings)

    with TestClient(application) as client:
        empty = client.post("/api/v1/search/export", json={"conditions": [], "logic": "and"})
        assert empty.status_code == 422
        invalid_operator = client.post(
            "/api/v1/search/export",
            json={
                "conditions": [
                    {"attribute_id": 1, "entry_id": 1, "operator": "like", "value": 650}
                ],
                "logic": "and",
            },
        )
        assert invalid_operator.status_code == 422
        unknown_entry = client.post(
            "/api/v1/search/export",
            json={
                "conditions": [
                    {"attribute_id": 1, "entry_id": 999, "operator": "gt", "value": 650}
                ],
                "logic": "and",
            },
        )
        assert unknown_entry.status_code == 202
        failed = _wait_for_job(client, unknown_entry.json()["job_id"])
        assert failed["status"] == "failed"
